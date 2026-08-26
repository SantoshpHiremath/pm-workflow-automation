"""
tests/test_excel_dashboard.py
-------------------------------

Verifies the generated Excel workbook is structurally correct AND that
the live formulas actually evaluate to the right numbers -- not just
that a file gets written. openpyxl reading a workbook it just wrote
only gives back the formula *string* by default, so to check real
evaluated values we convert to CSV via LibreOffice headless (the same
technique used manually during development) and parse the result.
"""

from __future__ import annotations

import csv
import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.feature_requests import SAMPLE_REQUESTS
from src.prioritization import score_all
from src.excel_dashboard import build_dashboard, COLUMNS

SOFFICE = shutil.which("soffice")


@pytest.fixture
def dashboard_path(tmp_path):
    scored = score_all(SAMPLE_REQUESTS)
    out = tmp_path / "dashboard.xlsx"
    build_dashboard(scored, out)
    return out


def test_build_dashboard_creates_a_file(dashboard_path):
    assert dashboard_path.exists()
    assert dashboard_path.stat().st_size > 0


def test_queue_sheet_has_expected_headers_and_row_count(dashboard_path):
    wb = load_workbook(dashboard_path)
    ws = wb["Prioritized Queue"]

    header_row = [cell.value for cell in ws[1]]
    assert header_row == COLUMNS

    # header row + one row per sample request
    assert ws.max_row == len(SAMPLE_REQUESTS) + 1


def test_rice_score_column_contains_a_live_formula_not_a_baked_value(dashboard_path):
    wb = load_workbook(dashboard_path)
    ws = wb["Prioritized Queue"]

    formula_cell = ws.cell(row=2, column=9).value
    assert isinstance(formula_cell, str)
    assert formula_cell.startswith("=ROUND(")


def test_summary_sheet_has_a_chart_and_tier_counts(dashboard_path):
    wb = load_workbook(dashboard_path)
    ws = wb["Summary"]

    assert len(ws._charts) == 1  # the BarChart was actually added

    tiers_seen = {ws.cell(row=r, column=1).value for r in range(2, 6)}
    assert tiers_seen == {"critical", "high", "medium", "low"}


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice (soffice) not installed in this environment")
def test_rice_formulas_evaluate_to_the_same_values_python_computed(dashboard_path, tmp_path):
    """The strongest check: actually evaluate the spreadsheet formulas
    (via LibreOffice headless, the same tool used to hand-verify this
    during development) and confirm they match compute_rice_score()."""
    scored = score_all(SAMPLE_REQUESTS)
    expected_scores = [s.rice_score for s in scored]

    csv_dir = tmp_path / "csv_out"
    csv_dir.mkdir()
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "csv", "--outdir", str(csv_dir), str(dashboard_path)],
        check=True, capture_output=True, timeout=60,
    )

    csv_path = csv_dir / (dashboard_path.stem + ".csv")
    with open(csv_path, newline="") as f:
        rows = list(csv.reader(f))

    # column index 8 (0-based) is "RICE Score"; row 0 is the header
    actual_scores = [float(row[8]) for row in rows[1:1 + len(expected_scores)]]
    assert actual_scores == expected_scores
