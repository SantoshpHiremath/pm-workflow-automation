"""
tests/test_workflow.py
------------------------

End-to-end test of the orchestration: intake -> score -> dashboard ->
notify. Runs the real scoring and real Excel-writing code, and uses
requests_mock only at the outermost HTTP boundary -- so this exercises
the actual integration between all four modules, not mocks calling
mocks.
"""

from __future__ import annotations

from openpyxl import load_workbook

from src.feature_requests import SAMPLE_REQUESTS
from src.workflow import run_pm_workflow

WEBHOOK_URL = "https://hooks.example.com/kinexon-pm-alerts"


def test_workflow_without_webhook_skips_notifications_entirely(tmp_path):
    result = run_pm_workflow(
        SAMPLE_REQUESTS, dashboard_output_path=tmp_path / "dashboard.xlsx", webhook_url=None,
    )

    assert result.notifications == []
    assert result.dashboard_path.exists()
    assert len(result.scored_requests) == len(SAMPLE_REQUESTS)


def test_workflow_with_webhook_notifies_only_critical_requests(tmp_path, requests_mock):
    requests_mock.post(WEBHOOK_URL, json={"ok": True}, status_code=200)

    result = run_pm_workflow(
        SAMPLE_REQUESTS, dashboard_output_path=tmp_path / "dashboard.xlsx", webhook_url=WEBHOOK_URL,
    )

    critical_ids = {s.request.request_id for s in result.scored_requests if s.priority_tier == "critical"}
    notified_ids = {n.request_id for n in result.notifications}
    assert notified_ids == critical_ids
    assert requests_mock.call_count == len(critical_ids)


def test_workflow_dashboard_file_is_a_real_readable_workbook_with_all_requests(tmp_path):
    result = run_pm_workflow(
        SAMPLE_REQUESTS, dashboard_output_path=tmp_path / "dashboard.xlsx", webhook_url=None,
    )

    wb = load_workbook(result.dashboard_path)
    ws = wb["Prioritized Queue"]
    assert ws.max_row == len(SAMPLE_REQUESTS) + 1

    written_ids = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    expected_ids = {r.request_id for r in SAMPLE_REQUESTS}
    assert written_ids == expected_ids


def test_workflow_result_scored_requests_are_sorted_critical_first(tmp_path):
    result = run_pm_workflow(
        SAMPLE_REQUESTS, dashboard_output_path=tmp_path / "dashboard.xlsx", webhook_url=None,
    )

    tier_order = {"critical": 0, "high": 1, "medium": 2, "low": 3}
    tiers = [tier_order[s.priority_tier] for s in result.scored_requests]
    assert tiers == sorted(tiers)
